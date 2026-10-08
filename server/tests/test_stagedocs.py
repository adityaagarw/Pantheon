"""Stage documents: spreadsheets, PDFs and notebooks shared by the user and agents."""

from __future__ import annotations

import io
import json

from openpyxl import Workbook, load_workbook

from app.models import Turn
from app.services import notebook_kernel
from tests.conftest import wait_for
from tests.helpers import call, dm, make_agent, make_org, rows, say, script, tool_results


def tiny_pdf(pages: list[str]) -> bytes:
    """A real (if minimal) PDF with one line of text per page."""
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", ""]
    kids = []
    for i, text in enumerate(pages):
        page, content = 3 + i * 2, 4 + i * 2
        kids.append(f"{page} 0 R")
        stream = f"BT /F1 18 Tf 72 720 Td ({text}) Tj ET"
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content} 0 R "
                    f"/Resources << /Font << /F1 {3 + len(pages) * 2} 0 R >> >> >>")
        objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    objs[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(pages)} >>"
    objs.append("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out, offsets = b"%PDF-1.4\n", []
    for n, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{n} 0 obj\n{body}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets).encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


async def _done(agent_id: str) -> bool:
    return any(t.status == "completed" for t in await rows(Turn, Turn.agent_id == agent_id))


async def test_spreadsheet_agent_and_user_work_together(app_client):
    org = await make_org(app_client)
    ana = await make_agent(app_client, org["id"], "Ana")
    prompts = script(
        ana["id"],
        call("sheet_new", title="Budget"),
        call("sheet_write", doc="Budget", start="A1",
             rows=[["Item", "Cost"], ["Rent", 1200], ["Food", "450"], ["Total", "=SUM(B2:B3)"]]),
        call("sheet_write", doc="Budget", cells=[{"cell": "C4", "formula": "B4/12"},
                                                 {"cell": "D1", "formula": "=1/0"}]),
        call("sheet_read", doc="Budget"),
        say("done"), then=say("ok"))
    await dm(app_client, org["id"], "Ana", "make a budget")
    await wait_for(lambda: _done(ana["id"]), timeout=30, msg="agent finished")
    results = tool_results(prompts[-1])
    assert "4 cells" in results[0] or "Created spreadsheet 'Budget'" in results[0]
    assert "Formula errors: D1 #DIV/0!" in results[2]
    table = results[3]
    assert "| 4 | Total | 1650 | 137.5 |" in table and "B4 =SUM(B2:B3) → 1650" in table

    doc = (await app_client.get(f"/api/v1/orgs/{org['id']}/docs")).json()[0]
    # The user edits in the browser; formulas that depend on it follow.
    # (The grid sends typed numbers as text; they're stored as numbers.)
    r = await app_client.put(f"/api/v1/docs/{doc['id']}/cells", json={
        "changes": {"Sheet1": {"B3": {"v": "550"}, "D1": None}}})
    assert r.status_code == 200
    full = (await app_client.get(f"/api/v1/docs/{doc['id']}")).json()
    cells = full["data"]["sheets"][0]["cells"]
    assert cells["B3"]["v"] == 550 and cells["B4"]["v"] == 1750 and "D1" not in cells

    # xlsx out and back in, formulas intact.
    x = await app_client.get(f"/api/v1/docs/{doc['id']}/export")
    wb = load_workbook(io.BytesIO(x.content))
    assert wb["Sheet1"]["B4"].value == "=SUM(B2:B3)"
    up = (await app_client.post(f"/api/v1/orgs/{org['id']}/attachments", files={
        "file": ("budget.xlsx", x.content, "application/octet-stream")})).json()
    again = (await app_client.post(f"/api/v1/orgs/{org['id']}/docs", json={
        "kind": "sheet", "attachmentId": up["id"]})).json()
    assert again["data"]["sheets"][0]["cells"]["B4"]["f"] == "=SUM(B2:B3)"
    assert again["data"]["sheets"][0]["cells"]["B4"]["v"] == 1750


async def test_csv_import():
    from app.services.stagedocs import sheets_from_file

    sheets = sheets_from_file("t.csv", b"name,score\nann,3\nben,4.5\n")
    assert sheets[0]["cells"]["B3"] == {"v": 4.5} and sheets[0]["cells"]["A1"] == {"v": "name"}
    wb = Workbook()
    wb.active["A1"] = 2
    wb.active["A2"] = "=A1*10"
    buf = io.BytesIO()
    wb.save(buf)
    cells = sheets_from_file("t.xlsx", buf.getvalue())[0]["cells"]
    assert cells["A2"]["f"] == "=A1*10"


async def test_pdf_reading_together(app_client):
    org = await make_org(app_client)
    leo = await make_agent(app_client, org["id"], "Leo")
    await app_client.post(f"/api/v1/orgs/{org['id']}/attachments", files={
        "file": ("report.pdf", tiny_pdf(["Introduction", "Revenue grew 12 percent"]),
                 "application/pdf")})
    prompts = script(
        leo["id"],
        call("pdf_open", file="report.pdf"),
        call("pdf_read", doc="report.pdf", page=2),
        call("pdf_show", doc="report.pdf", page=2, highlight="Revenue grew 12 percent",
             note="the key number"),
        say("look at page 2"), then=say("ok"))
    await dm(app_client, org["id"], "Leo", "walk me through the report")
    await wait_for(lambda: _done(leo["id"]), timeout=30, msg="agent finished")
    results = tool_results(prompts[-1])
    assert "Opened 'report.pdf'" in results[0]
    assert "Revenue grew 12 percent" in results[1] and "page 2 of 2" in results[1]
    doc = (await app_client.get(f"/api/v1/orgs/{org['id']}/docs")).json()[0]
    data = (await app_client.get(f"/api/v1/docs/{doc['id']}")).json()["data"]
    assert data["page"] == 2 and data["highlights"][0]["by"] == "Leo"
    pdf = await app_client.get(f"/api/v1/docs/{doc['id']}/export")
    assert pdf.content.startswith(b"%PDF")


async def test_notebook_cells_run_in_the_kernel(app_client, monkeypatch):
    ran: list[tuple[str, str]] = []

    async def fake_execute(notebook: str, code: str, timeout: float = 300):
        ran.append((notebook, code))
        if "boom" in code:
            return {"count": len(ran), "status": "error", "outputs": [
                {"type": "error", "ename": "NameError", "evalue": "name 'boom' is not defined",
                 "traceback": "NameError: name 'boom' is not defined"}]}
        return {"count": len(ran), "status": "ok", "outputs": [
            {"type": "stream", "name": "stdout", "text": "hello\n"},
            {"type": "image", "data": "iVBORw0KGgo="}, {"type": "result", "text": "42"}]}

    monkeypatch.setattr(notebook_kernel, "execute", fake_execute)
    org = await make_org(app_client)
    kai = await make_agent(app_client, org["id"], "Kai")
    prompts = script(
        kai["id"],
        call("notebook_new", title="Analysis"),
        call("notebook_add_cell", doc="Analysis", source="# Sales", type="markdown"),
        call("notebook_add_cell", doc="Analysis", source="print('hello'); 6*7", run=True),
        call("notebook_add_cell", doc="Analysis", source="boom", run=True),
        call("notebook_read", doc="Analysis"),
        say("done"), then=say("ok"))
    await dm(app_client, org["id"], "Kai", "analyse")
    await wait_for(lambda: _done(kai["id"]), timeout=30, msg="agent finished")
    results = tool_results(prompts[-1])
    assert "hello" in results[2] and "42" in results[2]
    assert "NameError" in results[3]
    assert "cell 1 (markdown)" in results[4] and "[a chart/image]" in results[4]
    ran_msgs = [m for m in prompts[-1] if m.type == "tool"]
    assert getattr(ran_msgs[2], "content", None) is not None

    doc = (await app_client.get(f"/api/v1/orgs/{org['id']}/docs")).json()[0]
    assert [c for _, c in ran] == ["print('hello'); 6*7", "boom"]
    assert all(nb == doc["id"] for nb, _ in ran)

    # .ipynb out and back in.
    nb = json.loads((await app_client.get(f"/api/v1/docs/{doc['id']}/export")).content)
    assert nb["nbformat"] == 4 and nb["cells"][0]["cell_type"] == "markdown"
    assert nb["cells"][1]["outputs"][1]["data"]["image/png"] == "iVBORw0KGgo="
    up = (await app_client.post(f"/api/v1/orgs/{org['id']}/attachments", files={
        "file": ("a.ipynb", json.dumps(nb).encode(), "application/json")})).json()
    again = (await app_client.post(f"/api/v1/orgs/{org['id']}/docs", json={
        "kind": "notebook", "attachmentId": up["id"]})).json()
    assert [c["type"] for c in again["data"]["cells"]] == ["markdown", "code", "code"]
    assert again["data"]["cells"][1]["outputs"][0]["text"] == "hello\n"

    # The user edits, moves and runs cells from the page.
    cid = again["data"]["cells"][2]["id"]
    await app_client.patch(f"/api/v1/docs/{again['id']}/cells/{cid}", json={"source": "1+1"})
    await app_client.patch(f"/api/v1/docs/{again['id']}/cells/{cid}", json={"to": 1})
    r = await app_client.post(f"/api/v1/docs/{again['id']}/cells/{cid}/run")
    assert r.status_code == 202
    await wait_for(lambda: _cell_ran(app_client, again["id"]), msg="user's cell ran")


async def _cell_ran(client, doc_id: str) -> bool:
    cells = (await client.get(f"/api/v1/docs/{doc_id}")).json()["data"]["cells"]
    return cells[0]["source"] == "1+1" and cells[0]["status"] == "ok"
