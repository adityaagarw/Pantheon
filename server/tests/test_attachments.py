"""Uploading documents and images to agents: extraction, search, tools, messages."""

from __future__ import annotations

import asyncio
import io
import struct
import zlib

import pytest
from langchain_core.messages import HumanMessage, ToolMessage

from app.models import Agent, Attachment, ToolCall
from tests.conftest import wait_for
from tests.helpers import call, dm, make_agent, make_org, rows, say, script


@pytest.fixture(autouse=True)
async def _drain_processing(app_client):
    """Uploads are processed in background tasks; let them finish before the DB is reset."""
    yield
    from app.services import attachments

    if attachments._tasks:
        await asyncio.gather(*list(attachments._tasks), return_exceptions=True)


def png(w: int = 40, h: int = 30, color=(200, 30, 30)) -> bytes:
    raw = b"".join(b"\x00" + bytes(color) * w for _ in range(h))

    def chunk(t: bytes, d: bytes) -> bytes:
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)

    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def tiny_pdf(text: str) -> bytes:
    stream = f"BT /F1 18 Tf 50 700 Td ({text}) Tj ET".encode()
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
            b"/Resources << /Font << /F1 5 0 R >> >> >>",
            b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return bytes(out)


def docx_bytes() -> bytes:
    from docx import Document

    d = Document()
    d.add_paragraph("The warranty period is 24 months from delivery.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Item", "Price"
    t.cell(1, 0).text, t.cell(1, 1).text = "Widget", "42"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def xlsx_bytes() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Region", "Revenue"])
    ws.append(["North", 1200])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def _upload(client, org_id: str, name: str, data: bytes, agent_id: str = "",
                  mime: str = "application/octet-stream") -> dict:
    r = await client.post(f"/api/v1/orgs/{org_id}/attachments",
                          files={"file": (name, data, mime)}, data={"agent_id": agent_id})
    assert r.status_code == 201, r.text
    return r.json()


async def _ready(att_id: str) -> Attachment | None:
    a = (await rows(Attachment, Attachment.id == att_id))[0]
    return a if a.status != "processing" else None


async def _done(agent_id: str, n: int) -> bool:
    calls = await rows(ToolCall, ToolCall.agent_id == agent_id)
    a = (await rows(Agent, Agent.id == agent_id))[0]
    return len(calls) >= n and a.runtime_status == "idle"


def _out(prompts, i: int) -> str:
    outs = [str(m.content) for m in prompts[i] if isinstance(m, ToolMessage)]
    return outs[-1] if outs else ""


async def test_documents_are_extracted_by_content_not_by_name(app_client):
    org = await make_org(app_client)
    ada = await make_agent(app_client, org["id"], "Ada")
    cases = {
        "report.pdf": (tiny_pdf("Quarterly revenue grew 12 percent"), "Quarterly revenue grew"),
        "contract.docx": (docx_bytes(), "warranty period is 24 months"),
        "sales.xlsx": (xlsx_bytes(), "North\t1200"),
        "notes.md": (b"# Plan\nShip the beta on Friday.\n", "Ship the beta"),
        "renamed-pdf.txt": (tiny_pdf("Really a PDF"), "Really a PDF"),  # sniffed, not trusted
    }
    for name, (data, expect) in cases.items():
        up = await _upload(app_client, org["id"], name, data, ada["id"])
        a = await wait_for(lambda att_id=up["id"]: _ready(att_id), msg=f"processed {name}")
        assert a.status == "ready", (name, a.error)
        text = (await app_client.get(f"/api/v1/attachments/{up['id']}/text")).json()["text"]
        assert expect in text, (name, text[:200])
    pdf = (await rows(Attachment, Attachment.name == "renamed-pdf.txt"))[0]
    assert pdf.mime == "application/pdf" and pdf.info["pages"] == 1
    listing = (await app_client.get(f"/api/v1/orgs/{org['id']}/attachments",
                                    params={"agent_id": ada["id"]})).json()
    assert len(listing) == 5


async def test_bad_uploads_are_rejected_or_marked_failed(app_client):
    org = await make_org(app_client)
    empty = await app_client.post(f"/api/v1/orgs/{org['id']}/attachments",
                                  files={"file": ("a.txt", b"", "text/plain")})
    assert empty.status_code == 400
    up = await _upload(app_client, org["id"], "broken.pdf", b"%PDF-1.4 this is not a real pdf")
    a = await wait_for(lambda: _ready(up["id"]), msg="processed")
    assert a.status == "failed" and "couldn't read" in a.error
    ghost = await app_client.post(f"/api/v1/orgs/{org['id']}/attachments",
                                  files={"file": ("a.txt", b"hi", "text/plain")},
                                  data={"agent_id": "agt_missing"})
    assert ghost.status_code == 404
    blob = await _upload(app_client, org["id"], "data.bin", bytes(range(256)) * 4)
    assert blob["kind"] == "other"


async def test_agents_search_read_and_see_files_in_their_prompt(app_client):
    org = await make_org(app_client)
    ada = await make_agent(app_client, org["id"], "Ada")
    ben = await make_agent(app_client, org["id"], "Ben")
    filler = "\n\n".join(f"Section {i}: unrelated housekeeping detail number {i}." for i in range(40))
    clause = " ".join(["The warranty period is 24 months from delivery, per clause 9."] * 10)
    doc = filler + "\n\n" + clause + "\n\n" + filler
    mine = await _upload(app_client, org["id"], "terms.txt", doc.encode(), ada["id"], "text/plain")
    shared = await _upload(app_client, org["id"], "handbook.md", b"Vacation policy: 25 days.\n")
    await wait_for(lambda: _ready(mine["id"]), msg="mine")
    await wait_for(lambda: _ready(shared["id"]), msg="shared")

    prompts = script(ada["id"], call("attachment_list"),
                     call("attachment_search", query="how long is the warranty?"),
                     call("attachment_read", file="terms.txt", offset=0, length=300),
                     say("done"))
    await dm(app_client, org["id"], "Ada", "Warranty period: how many months from delivery?")
    await wait_for(lambda: _done(ada["id"], 3), msg="tools")
    system = str(prompts[0][0].content)
    assert "# Files you've been given" in system and "terms.txt" in system
    assert "handbook.md" in system and "shared with the org" in system
    assert "24 months" in system  # the relevant passage was recalled into the turn
    assert "handbook.md" in _out(prompts, 1) and "terms.txt" in _out(prompts, 1)
    assert "24 months from delivery" in _out(prompts, 2)
    assert "continues: call again with offset=" in _out(prompts, 3)

    # Ben has his own library plus the shared file, but not Ada's.
    bp = script(ben["id"], call("attachment_list"), say("ok"))
    await dm(app_client, org["id"], "Ben", "what files do I have?")
    await wait_for(lambda: _done(ben["id"], 1), msg="ben")
    assert "handbook.md" in _out(bp, 1) and "terms.txt" not in _out(bp, 1)


async def test_images_reach_vision_models_only_and_download_safely(app_client):
    org = await make_org(app_client)
    seer = await make_agent(app_client, org["id"], "Seer", model={"vision": True})
    blind = await make_agent(app_client, org["id"], "Blind")
    img = await _upload(app_client, org["id"], "chart.png", png(), seer["id"], "image/png")
    assert img["kind"] == "image"
    a = await wait_for(lambda: _ready(img["id"]), msg="image")
    assert a.info == {"width": 40, "height": 30}

    up = await app_client.get(f"/api/v1/attachments/{img['id']}/file")
    assert up.status_code == 200 and up.headers["content-type"] == "image/png"
    assert up.headers["x-content-type-options"] == "nosniff"
    assert "sandbox" in up.headers["content-security-policy"]
    html = await _upload(app_client, org["id"], "evil.html",
                         b"<script>alert(1)</script>", seer["id"], "text/html")
    r = await app_client.get(f"/api/v1/attachments/{html['id']}/file")
    assert r.headers["content-type"] == "application/octet-stream"
    assert r.headers["content-disposition"].startswith("attachment")

    sp = script(seer["id"], call("attachment_view", file="chart.png"), say("A red block."))
    await dm(app_client, org["id"], "Seer", "Look at the chart.")
    await wait_for(lambda: _done(seer["id"], 1), msg="view")
    assert any(isinstance(m, HumanMessage) and isinstance(m.content, list) for m in sp[1])

    bp = script(blind["id"], say("ok"))
    shared = await _upload(app_client, org["id"], "shared.png", png())
    await wait_for(lambda: _ready(shared["id"]), msg="shared image")
    tool = script(blind["id"], call("attachment_view", file="shared.png"), say("cannot see"))
    await dm(app_client, org["id"], "Blind", "Look at shared.png")
    await wait_for(lambda: _done(blind["id"], 1), msg="blind")
    assert "isn't set up to see images" in _out(tool, 1)
    _ = bp


async def test_attaching_files_to_a_message(app_client):
    org = await make_org(app_client)
    seer = await make_agent(app_client, org["id"], "Seer", model={"vision": True})
    doc = await _upload(app_client, org["id"], "brief.txt", b"Launch is on the 14th.", seer["id"])
    img = await _upload(app_client, org["id"], "photo.png", png(), seer["id"])
    await wait_for(lambda: _ready(doc["id"]), msg="doc")
    await wait_for(lambda: _ready(img["id"]), msg="img")

    prompts = script(seer["id"], say("Got both."))
    r = await app_client.post(f"/api/v1/orgs/{org['id']}/messages", json={
        "to": "Seer", "content": "Here you go.", "attachments": [doc["id"], img["id"]]})
    assert r.status_code == 201
    msg = r.json()
    assert [a["name"] for a in msg["meta"]["attachments"]] == ["brief.txt", "photo.png"]
    await wait_for(lambda: len(prompts) >= 1, msg="turn")
    human = [m for m in prompts[0] if isinstance(m, HumanMessage)]
    text = str(human[0].content)
    assert "[Attached: brief.txt (document" in text and "photo.png (image" in text
    assert any(isinstance(m.content, list) for m in human)  # the image itself, for vision

    only = await app_client.post(f"/api/v1/orgs/{org['id']}/messages",
                                 json={"to": "Seer", "attachments": [doc["id"]]})
    assert only.status_code == 201 and "attached 1 file" in only.json()["content"]
    other = await make_org(app_client, name="Other Co")
    bad = await app_client.post(f"/api/v1/orgs/{other['id']}/messages",
                                json={"to": "Nobody", "content": "hi", "attachments": [doc["id"]]})
    assert bad.status_code == 400  # files from another organization can't be attached


async def test_deleting_files_agents_and_orgs_cleans_up(app_client):
    from app.core.config import settings

    org = await make_org(app_client)
    ada = await make_agent(app_client, org["id"], "Ada")
    one = await _upload(app_client, org["id"], "one.txt", b"first file", ada["id"])
    two = await _upload(app_client, org["id"], "two.txt", b"second file", ada["id"])
    await wait_for(lambda: _ready(one["id"]), msg="one")
    await wait_for(lambda: _ready(two["id"]), msg="two")
    folder = settings.uploads_path / org["id"] / one["id"]
    assert (folder / "file").is_file()
    assert (await app_client.delete(f"/api/v1/attachments/{one['id']}")).status_code == 204
    assert not folder.exists()
    assert (await app_client.get(f"/api/v1/attachments/{one['id']}")).status_code == 404

    assert (await app_client.delete(f"/api/v1/agents/{ada['id']}")).status_code == 204
    assert not await rows(Attachment, Attachment.id == two["id"])
    three = await _upload(app_client, org["id"], "three.txt", b"third")
    await wait_for(lambda: _ready(three["id"]), msg="three")
    assert (await app_client.delete(f"/api/v1/orgs/{org['id']}")).status_code == 204
    assert not (settings.uploads_path / org["id"]).exists()
