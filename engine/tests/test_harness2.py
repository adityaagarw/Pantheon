"""What Pantheon took from harness2 (Google Research, 2026): blocked hosts, reopening
deliverables, skills, tool-loop notes, the checker's criteria, the experience library,
rubric-graded benchmarks and team tuning."""

from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from langchain_core.messages import AIMessage, HumanMessage, ToolMessage

from app.agents import nudges
from app.agents.prompt import build_system_prompt
from app.core import egress
from app.models import Agent, Org
from app.services import (  # noqa: F401 (bench: swarmbench on sys.path)
    acceptance,
    bench,
    experience,
    filecheck,
    skills,
    tuning,
)
from app.tools.base import ToolContext, ToolError
from tests.helpers import make_agent, make_org, rows


def _org(**settings) -> SimpleNamespace:
    return SimpleNamespace(id="org_x", settings=settings)


# --- blocked hosts ----------------------------------------------------------------------------


def test_blocked_hosts_cover_subdomains_and_path_prefixes():
    entries = egress.hosts(_org(blocked_hosts=["huggingface.co", "https://github.com/Job-Bench/"]))
    assert entries == ("huggingface.co", "github.com/job-bench")
    assert egress.blocked_url("https://huggingface.co/datasets/JobBench/x", entries) == "huggingface.co"
    assert egress.blocked_url("https://cdn-lfs.huggingface.co/a", entries) == "huggingface.co"
    assert egress.blocked_url("https://github.com/Job-Bench/job-bench-eval", entries)
    assert egress.blocked_url("https://github.com/other/repo", entries) is None
    assert egress.blocked_url("https://nothuggingface.co/", entries) is None
    org = _org(blocked_hosts="huggingface.co")
    assert egress.blocked_command(org, "curl -L https://huggingface.co/datasets/x") == "huggingface.co"
    assert egress.blocked_command(org, "curl https://mynothuggingface.co/") is None
    assert egress.blocked_command(_org(), "curl https://huggingface.co") is None


async def test_fetch_url_refuses_a_blocked_host():
    from app.tools.web import fetch_url

    ctx = SimpleNamespace(org=_org(blocked_hosts=["huggingface.co"]), agent=SimpleNamespace(id="a"))
    with pytest.raises(ToolError, match="blocked for this organization"):
        await fetch_url({"url": "https://huggingface.co/datasets/JobBench/job-bench"}, ctx)


# --- reopening deliverables -------------------------------------------------------------------


def _workbook(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Recon"
    ws.append(["Item", "Amount"])
    ws.append(["Rent", 9600])
    ws.append(["Total", "=SUM(B2:B2)"])  # openpyxl saves the formula and no value
    ws.append(["Broken", "#REF!"])
    wb.create_sheet("Empty")
    wb.save(path)


def test_a_workbook_is_read_as_a_program_reads_it(tmp_path):
    _workbook(tmp_path / "book.xlsx")
    r = filecheck.check_file(tmp_path / "book.xlsx", "book.xlsx")
    text = "; ".join(r.problems)
    assert "1 formula cell has no saved value (e.g. Recon!B3 =SUM(B2:B2))" in text
    assert "1 cell shows an error: Recon!B4 #REF!" in text and "sheet 'Empty' is empty" in text
    assert r.summary.startswith("2 sheets: Recon")


def test_csv_documents_and_pages_show_what_breaks_their_reader(tmp_path):
    (tmp_path / "log.csv").write_text("ID,Amount,Note\n1,10,ok\n2,1,000,shifted\n", encoding="utf-8")
    r = filecheck.check_file(tmp_path / "log.csv", "log.csv")
    assert "1 row has a different number of fields from the header's 3 (line 3)" in r.problems[0]

    import docx

    d = docx.Document()
    d.add_paragraph("Memo dated [insert date] for the owner. TODO: totals.")
    d.save(tmp_path / "memo.docx")
    r = filecheck.check_file(tmp_path / "memo.docx", "memo.docx")
    assert "placeholder text left in" in r.problems[0] and "[insert date]" in r.problems[0]

    (tmp_path / "index.html").write_text('<a href="about.html">About</a><img src="https://x/y.png">'
                                         '<a href="#top">Top</a>', encoding="utf-8")
    r = filecheck.check_file(tmp_path / "index.html", "index.html")
    assert r.problems == ["links to files that don't exist: about.html"]
    (tmp_path / "empty.json").write_text("", encoding="utf-8")
    assert filecheck.check_file(tmp_path / "empty.json", "empty.json").problems == [
        "the file is empty (0 bytes)"]


def test_the_files_a_request_names_and_the_ones_missing(tmp_path):
    request = ('Deliver three files. First, "Q1_2024_Reconciliation_Workbook.xlsx" with four '
               'sheets. Second, "Q1 Expense Summary.csv". You need the external '
               'Accounts-Payable.xlsx template. Local files:\n- Bank_Statement_Q1_2024.csv')
    names = filecheck.expected_files(request)
    assert names == ["Q1_2024_Reconciliation_Workbook.xlsx", "Q1 Expense Summary.csv",
                     "Accounts-Payable.xlsx", "Bank_Statement_Q1_2024.csv"]
    (tmp_path / "output").mkdir()
    (tmp_path / "Bank_Statement_Q1_2024.csv").write_text("a\n1\n", encoding="utf-8")
    (tmp_path / "output" / "q1 expense summary.csv").write_text("a\n1\n", encoding="utf-8")
    assert filecheck.missing_files(tmp_path, names) == ["Q1_2024_Reconciliation_Workbook.xlsx",
                                                        "Accounts-Payable.xlsx"]


def test_a_write_reports_only_problems_in_what_it_wrote(tmp_path):
    (tmp_path / "ok.csv").write_text("a,b\n1,2\n", encoding="utf-8")
    assert filecheck.after_tool_note(tmp_path, "write_file", {"path": "ok.csv"}, 0) == ""
    (tmp_path / "bad.csv").write_text("a,b\n1,2,3\n", encoding="utf-8")
    note = filecheck.after_tool_note(tmp_path, "write_file", {"path": "bad.csv"}, 0)
    assert note.startswith("\n\n[Pantheon reopened the files this produced") and "bad.csv" in note
    _workbook(tmp_path / "made.xlsx")  # as if a script run_command started had written it
    note = filecheck.after_tool_note(tmp_path, "run_command", {"command": "python make.py"}, 0)
    assert "- made.xlsx: " in note and "1 formula cell has no saved value" in note
    assert "ok.csv" not in note


async def test_the_checker_is_told_what_reopening_the_files_found(tmp_path):
    (tmp_path / "out.csv").write_text("a,b\n1\n", encoding="utf-8")
    note = acceptance._files_note(tmp_path, 'Write "out.csv" and "summary.docx".')
    assert "- out.csv: 1 row has a different number of fields" in note
    assert "The request names summary.docx, and no file of that name" in note
    assert "{files}" in acceptance.FIRST and "Values are values" in acceptance.FIRST
    assert "check_files" in acceptance.LOOK


# --- skills -----------------------------------------------------------------------------------


def test_a_skill_says_when_to_use_it_and_ships_its_helpers():
    with pytest.raises(skills.SkillError, match="must say when to use"):
        skills.save("contract-review", "Contract review know-how.", "Read every clause.")
    with pytest.raises(skills.SkillError, match="template text left"):
        skills.save("contract-review", "Use when reviewing contracts.", "Fill {{client}} in.")
    with pytest.raises(skills.SkillError, match="doesn't include"):
        skills.save("contract-review", "Use when reviewing contracts.", "Run scripts/dates.py.")
    s = skills.save("contract-review", "Use when reviewing contracts.", "Run scripts/dates.py.",
                    {"scripts/dates.py": "print('ok')"})
    assert s.files == ["scripts/dates.py"] and not s.builtin
    assert {"spreadsheet-deliverables", "source-reconciliation", "contract-review"} <= set(
        skills.all_skills())
    section, missing = skills.prompt_section(["contract-review", "nope"])
    assert "- contract-review: Use when reviewing contracts." in section
    assert missing == ["skill 'nope' doesn't exist"]
    skills.delete("contract-review")
    with pytest.raises(skills.SkillError):
        skills.delete("spreadsheet-deliverables")  # built in


async def test_an_agent_sees_its_skills_triggers_and_opens_one(app_client):
    from app.tools.skilltools import use_skill

    skills.save("ledger-close", "Use when closing a month's books.", "Tie out scripts/tie.py.",
                {"scripts/tie.py": "print('tied')"})
    org = await make_org(app_client)
    ann = await make_agent(app_client, org["id"], "Ann", skills=["ledger-close"])
    o = (await rows(Org, Org.id == org["id"]))[0]
    a = (await rows(Agent, Agent.id == ann["id"]))[0]
    prompt = await build_system_prompt(o, a)
    assert "# Skills" in prompt and "- ledger-close: Use when closing a month's books." in prompt
    assert "Tie out" not in prompt  # the instructions stay out until opened
    from app.tools.resolve import resolve_tools

    assert (await resolve_tools(o, a)).tools["use_skill"].core
    ctx = ToolContext(org=o, agent=a, turn_id=None, tool_call_id="t1")
    text = await use_skill({"name": "ledger-close"}, ctx)
    assert text.startswith("# Skill: ledger-close\nTie out scripts/tie.py.")
    assert (Path(ctx.workdir) / ".pantheon/skills/ledger-close/scripts/tie.py").is_file()
    with pytest.raises(ToolError, match="no skill named"):
        await use_skill({"name": "spreadsheet-deliverables"}, ctx)

    # An organization file carries the user's own skills, so it works elsewhere.
    from app.services import orgs

    definition = await orgs.export_org(org["id"])
    assert definition["skillDefs"]["ledger-close"]["files"] == {"scripts/tie.py": "print('tied')"}
    skills.delete("ledger-close")
    await orgs.import_org(definition, name="Copy")
    assert skills.get("ledger-close") is not None
    skills.delete("ledger-close")


# --- notes in the tool loop -------------------------------------------------------------------


def _calls(n: int, name: str = "read_file", result: str = "x") -> list:
    out = []
    for i in range(n):
        out += [AIMessage("", tool_calls=[{"id": f"{name}{i}", "name": name, "args": {}}]),
                ToolMessage(result, tool_call_id=f"{name}{i}")]
    return out


async def test_a_long_stretch_without_writing_gets_a_note_once(app_client):
    org = await make_org(app_client)
    ann = await make_agent(app_client, org["id"], "Ann")
    await app_client.post(f"/api/v1/orgs/{org['id']}/tasks",
                          json={"title": "Reconcile Q1", "assigneeId": ann["id"]})
    agent = (await rows(Agent, Agent.id == ann["id"]))[0]
    msgs = [HumanMessage("go"), *_calls(30)]
    lines, noted = await nudges.notes(agent, msgs, 30, {"todo": True})
    assert any("haven't written any of the work yet" in x for x in lines)
    again, _ = await nudges.notes(agent, msgs, 31, noted)
    assert not any("haven't written" in x for x in again)
    wrote = [*msgs, *_calls(1, "write_file")]
    lines, _ = await nudges.notes(agent, wrote, 31, {"todo": True})
    assert not any("haven't written" in x for x in lines)


async def test_the_same_failure_three_times_gets_a_note():
    agent = SimpleNamespace(id="agt_x", org_id="org_none", todos=[1], is_supervisor=False)
    msgs = [HumanMessage("go"), *_calls(3, "sheet_write", "ERROR: invalid arguments at rows")]
    assert nudges.repeated_error(msgs) == ("sheet_write", "ERROR: invalid arguments at rows")
    lines, noted = await nudges.notes(agent, msgs, 3, {})
    assert any("sheet_write failed the same way 3 times" in x for x in lines)
    again, _ = await nudges.notes(agent, msgs, 4, noted)
    assert not any("failed the same way" in x for x in again)
    mixed = [HumanMessage("go"), *_calls(2, "sheet_write", "ERROR: a"), *_calls(1, "x", "ERROR: b")]
    assert nudges.repeated_error(mixed) is None


# --- the experience library -------------------------------------------------------------------


LESSON = {"cue": "a reconciliation with a subledger", "lesson": "inventory every input first",
          "evidence": "run B opened the AP file and found the partial payment", "confidence": "medium"}


def test_the_library_keeps_lessons_and_refuses_a_graders_words():
    experience.add("bookkeeping", "jb-task1", team={"agents": [{"name": "Ann", "role": "Clerk",
                                                                "persona": "You...", "skills": []}]},
                   why="B delivered all three files.", lesson=LESSON,
                   edits=[{"summary": "inv1 add source-reconciliation", "verdict": "helped"},
                          {"summary": "dom1 persona", "verdict": "made up"}])
    assert experience.domains() == ["bookkeeping"]
    index = experience.read("bookkeeping")
    assert "- cue: a reconciliation with a subledger" in index
    entry = experience.read("bookkeeping", "jb-task1")
    assert "- [helped] inv1 add source-reconciliation" in entry
    assert "- [confounded] dom1 persona" in entry and "- Ann (Clerk)" in entry
    assert "inventory every input first" in experience.prompt_index()
    with pytest.raises(experience.ExperienceError, match="names grading"):
        experience.add("bookkeeping", "x", team={}, why="The rubric wanted CHK-1402.",
                       lesson=LESSON, edits=[])
    with pytest.raises(experience.ExperienceError, match="needs a lesson"):
        experience.add("bookkeeping", "x", team={}, why="", lesson={"cue": "c"}, edits=[])


# --- rubric-graded benchmarks -----------------------------------------------------------------


def _rubric_task(tmp_path: Path):
    from swarmbench import external, tasks

    rubric = {"rubrics": [
        {"rubric": "Is the total right?", "weight": 10, "criterion": ["Total is 9600", "Labelled"]},
        {"rubric": "Is there a memo?", "weight": 5, "criterion": "A memo exists"}]}
    external.write_task(tmp_path, "jb-clerk-task1", title="Clerk task 1", category="clerks",
                        summary="Reconcile.", prompt='Deliver "Recon.xlsx".',
                        files={"ledger.csv": b"a,b\n1,2\n"}, rubric=rubric,
                        source={"benchmark": "JobBench"}, blocked_hosts=["huggingface.co"])
    return tasks.load_tasks("jb-clerk-task1", tasks_dir=tmp_path)[0]


def test_an_imported_task_keeps_its_rubric_out_of_the_work(tmp_path):
    from swarmbench import prompt

    task = _rubric_task(tmp_path)
    assert task.grading == "rubric" and task.output_dir == "output"
    assert task.weights["rubric"] == 90 and task.weights["hidden_tests"] == 0
    assert not (task.repo_path / "rubric").exists() and task.rubric_path.is_file()
    text = prompt.build_prompt(task, None, "/w")
    assert "Write every deliverable to `output/`" in text and "hidden test" not in text


def test_a_judge_grades_what_a_program_reads(tmp_path):
    from swarmbench import grader, rubric, utils

    task = _rubric_task(tmp_path / "tasks")
    work = tmp_path / "work"
    utils.copy_tree(task.repo_path, work)
    (work / "output").mkdir()
    _workbook(work / "output" / "Recon.xlsx")
    seen = []

    def judge(p: str) -> str:
        seen.append(p)
        if "Is the total right?" in p:
            return '<think>hm</think>{"met": [false, true], "why": "B3 is blank"}'
        return '{"met": [true], "why": "ok"}'

    g = grader.grade_task(task, work, None, judge)
    r = g["parts"]["rubric"]
    assert r["score"] == round((10 * 0.5 + 5) / 15, 4) and r["items_met"] == 1
    assert r["files"] == ["output/Recon.xlsx"]
    shown = next(p for p in seen if "Is the total right?" in p)
    assert "\nRent\t9600\nTotal\n" in shown and "=SUM" not in shown  # saved values, as pandas reads
    assert g["weights"] == {"rubric": 90, "safety": 10} and g["score"] == round(
        (90 * r["score"] + 10) / 100 * 100, 1)
    assert rubric.load_rubric(task)[1]["criteria"] == ["A memo exists"]
    nothing = grader.grade_task(task, tmp_path / "missing", None, judge)
    assert nothing["parts"]["rubric"]["note"] == "nothing was delivered"


def test_lab_criteria_become_rubric_items():
    from swarmbench import external

    r = external.lab_rubric([{"id": "C-001", "title": "Pinnacle named", "match_criteria": "PASS if"}])
    assert r == {"rubrics": [{"rubric": "Pinnacle named", "weight": 1, "criterion": ["PASS if"]}]}
    assert "documents/" in external.lab_prompt({"instructions": "Draft it.",
                                                "deliverables": {"a.docx": "the draft"}})


# --- team tuning ------------------------------------------------------------------------------


BASE = {"name": "Clerks", "description": "Books.", "agents": [
    {"name": "Ann", "role": "Lead", "persona": "You lead.", "tools": [{"name": "read_file"}],
     "skills": []},
    {"name": "Bo", "role": "Clerk", "persona": "You match.", "tools": ["read_file"], "skills": []}],
    "relationships": [{"from": "Ann", "to": "Bo", "kind": "manages"}]}


def test_edits_apply_to_a_copy_of_the_team():
    edits = tuning.clean_edits([
        {"kind": "persona", "agent": "Bo", "append": "Open every input first.", "secret": 1},
        {"kind": "tools", "agent": "all", "add": ["check_files", "not_a_tool"], "remove": ["read_file"]},
        {"kind": "skill", "agent": "Ann", "name": "recon", "description": "Use when...", "body": "b"},
        {"kind": "description", "append": "One owner per file."},
        {"kind": "hire", "agent": "Cy"}], "inv")
    assert [e["id"] for e in edits] == ["inv1", "inv2", "inv3", "inv4"] and "secret" not in edits[0]
    out, problems = tuning.apply_edits(BASE, edits, {"recon": "recon-2"})
    assert problems == [] and BASE["agents"][1]["persona"] == "You match."
    bo = out["agents"][1]
    assert bo["persona"] == "You match.\n\nOpen every input first."
    assert [t["name"] for t in bo["tools"]] == ["check_files"]
    assert out["agents"][0]["skills"] == ["recon-2"] and out["description"].endswith("One owner per file.")
    _, problems = tuning.apply_edits(BASE, [{"id": "x1", "kind": "persona", "agent": "Zed"}])
    assert problems == ["x1: no agent 'Zed'"]


async def test_a_tuning_probes_both_edits_composes_and_applies(app_client, monkeypatch):
    org = await make_org(app_client, "Clerks")
    await make_agent(app_client, org["id"], "Ann", persona="You lead.")

    monkeypatch.setattr(bench, "check_models", lambda org: _none())
    runs: list[str] = []
    import asyncio

    gate = asyncio.Event()

    async def attempt(t, label, definition, task, wf, minutes, org):
        await gate.wait()  # held until the test has tried a second tuning
        runs.append(label)
        t.runs[label] = {"status": "done", "score": {"base": 50, "inv": 70, "dom": 55,
                                                     "final": 72}[label]}
        persona = definition["agents"][0]["persona"]
        return f"record of {label}: persona {persona!r}", {"score": t.runs[label]["score"]}

    async def propose(scope, task_text, definition, rec):
        assert "score" not in rec.lower() and "record of base" in rec
        if scope == "inv":
            return tuning.clean_edits([{"kind": "persona", "agent": "Ann",
                                        "append": "Inventory inputs.", "summary": "inventory"}], "inv")
        return tuning.clean_edits([{"kind": "skill", "agent": "Ann", "name": "recon-practice",
                                    "description": "Use when reconciling.", "body": "Match strictly.",
                                    "summary": "a skill"}], "dom")

    composed = {}

    async def compose(task_text, versions):
        composed["labels"] = [v[0] for v in versions]
        return {"adopt": ["inv1", "dom1"], "verdicts": {"inv1": "helped", "dom1": "confounded"},
                "why": "Both versions delivered more.", "lesson": LESSON}

    monkeypatch.setattr(tuning, "_attempt", attempt)
    monkeypatch.setattr(tuning, "propose", propose)
    monkeypatch.setattr(tuning, "compose", compose)
    t = await tuning.start(org["id"], "month-end-close")
    with pytest.raises(tuning.TuningError, match="already running"):
        await tuning.start(org["id"], "month-end-close")
    gate.set()
    from tests.conftest import wait_for

    await wait_for(lambda: t.status != "running", msg="tuning")
    assert t.status == "done", t.detail
    assert runs == ["base", "inv", "dom", "final"] and composed["labels"] == ["base", "inv", "dom"]
    assert t.adopted == ["inv1", "dom1"] and t.skills_made == ["recon-practice"]
    assert "finance-ops" in experience.domains() and t.library_entry
    saved = json.loads((tuning.folder() / f"{t.id}.json").read_text(encoding="utf-8"))
    assert saved["runs"]["final"]["score"] == 72

    changed = await tuning.apply(t.id)
    assert changed == ["Ann: persona, skills"]
    a = (await rows(Agent, Agent.org_id == org["id"], Agent.name == "Ann"))[0]
    assert a.persona.endswith("Inventory inputs.") and a.skills == ["recon-practice"]
    skills.delete("recon-practice")


async def test_cancel_stops_the_run_in_progress_and_grades_nothing(app_client, monkeypatch):
    import asyncio

    from tests.conftest import wait_for

    org = await make_org(app_client, "Clerks two")
    await make_agent(app_client, org["id"], "Bo", persona="You lead.")
    monkeypatch.setattr(bench, "check_models", lambda org: _none())
    monkeypatch.setattr(bench, "_fresh_folder", lambda src, dest: None)
    started = asyncio.Event()

    async def run_attempt(att, *args, **kwargs):
        started.set()
        while not att.cancel:
            await asyncio.sleep(0.01)
        att.status = "cancelled"
        return att

    async def grade_folder(*args):
        raise AssertionError("a cancelled run is not graded")

    monkeypatch.setattr(bench, "run_attempt", run_attempt)
    monkeypatch.setattr(bench, "grade_folder", grade_folder)
    t = await tuning.start(org["id"], "month-end-close")
    await asyncio.wait_for(started.wait(), 5)
    await tuning.cancel(t.id)
    await wait_for(lambda: t.status != "running", msg="tuning")
    assert t.status == "cancelled" and t.runs["base"]["status"] == "cancelled"
    assert not any(a.label.startswith(f"tune-{t.id}") for a in bench.ATTEMPTS.values())

    # One left "running" on disk when Pantheon stopped reads back as interrupted.
    t.status = "running"
    tuning._save(t)
    tuning.TUNINGS.pop(t.id)
    assert tuning.load(t.id).status == "interrupted"


def test_a_verdict_at_the_top_keeps_the_problems_under_it():
    report = ("**VERDICT: FAIL — the memo is missing.**\n\n1. **No memo.** doc_list shows only the sheet.\n"
              "2. Row 9 has no risk.\n\nCoverage: 12 rules, 10 tested.\n\nVERDICT: FAIL\n\n"
              "[subagent: 14 steps]")
    problems = acceptance._problems(report)
    assert acceptance.verdict(report) is False
    assert problems.startswith("1. **No memo.**") and "2. Row 9" in problems and "Coverage" in problems
    assert "VERDICT" not in problems and "subagent" not in problems


def test_the_checker_can_find_stage_documents():
    assert "doc_list" in acceptance.LOOK and "doc_list lists them" in acceptance.FIRST


async def _none() -> None:
    return None
