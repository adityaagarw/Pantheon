"""Deterministic checks of the files a team delivers: each file reopened the way its reader
will open it.

From harness2 (Google Research, 2026): its best case study (43.1 -> 77.8) added, among other
edits, a script that reopened each produced file, and found what agents never saw. A
workbook written by a script keeps its formulas but no saved values, so pandas, previews
and graders read those cells as blank; a CSV row with a field too many shifts every column
after it; a document says "[insert date]".
None of that shows in the tool result that wrote the file. So:

- check_files (a tool) shows an agent what a reader sees in its files, with the problems;
- after write_file, edit_file and run_command, the files they wrote are checked, and only
  problems are added to the tool's result;
- the acceptance check gets the same findings, and every file the request names that isn't
  in the workspace, before its checker starts.
"""

from __future__ import annotations

import csv
import io
import json
import os
import re
from pathlib import Path

CHECKED = {".xlsx", ".xlsm", ".csv", ".tsv", ".docx", ".pdf", ".json", ".md", ".txt",
           ".html", ".htm"}
OFFICE = {".xlsx", ".xlsm", ".csv", ".tsv", ".docx", ".pdf", ".json"}
SKIP = {".acceptance", ".pantheon"}  # the checker's and Pantheon's own folders
MAX_BYTES = 50_000_000
MAX_CELLS = 200_000
MAX_FILES = 40
WALK_CAP = 5000  # directory entries a scan for changed files looks at
ERRORS = ("#REF!", "#NAME?", "#DIV/0!", "#VALUE!", "#N/A", "#NUM!", "#NULL!", "#SPILL!")
PLACEHOLDER = re.compile(
    r"\b(?:TODO|TBD|FIXME|XXX)\b|lorem ipsum|\[(?:insert|add|your|enter|placeholder)\b[^\]]{0,40}\]"
    r"|\{\{\s*\w[^}]{0,40}\}\}|<(?:insert|placeholder)[^>]{0,40}>", re.I)
# A file the request names: "Q1 Reconciliation.xlsx", `report.md`, output/summary.csv.
_EXT = r"\.(?:xlsx|xlsm|csv|tsv|docx|pdf|json|md|txt|html|pptx)"
QUOTED = re.compile(r"[\"'`“‘]([^\"'`”’\n]{1,120}?" + _EXT + r")[\"'`”’]",
                    re.I)
BARE = re.compile(r"(?<![\w./\\-])([\w./\\-]{1,120}" + _EXT + r")(?![\w-])", re.I)


class Result:
    def __init__(self, rel: str) -> None:
        self.rel = rel
        self.summary = ""
        self.problems: list[str] = []

    def line(self, *, problems_only: bool = False) -> str:
        if problems_only:
            return f"- {self.rel}: " + "; ".join(self.problems)
        head = f"- {self.rel}: {self.summary or 'checked'}"
        return head + "".join(f"\n  - PROBLEM: {p}" for p in self.problems)


def _cell(ws_title: str, cell) -> str:  # noqa: ANN001 - an openpyxl cell
    return f"{ws_title}!{cell.coordinate}"


def _xlsx(path: Path, r: Result) -> None:
    from openpyxl import load_workbook

    wb_f = load_workbook(path, read_only=True)
    wb_v = load_workbook(path, read_only=True, data_only=True)
    sheets, blank, errors, cells = [], [], [], 0
    for ws_f, ws_v in zip(wb_f.worksheets, wb_v.worksheets, strict=False):
        rows = filled = 0
        for row_f, row_v in zip(ws_f.iter_rows(), ws_v.iter_rows(), strict=False):
            rows += 1
            for c_f, c_v in zip(row_f, row_v, strict=False):
                cells += 1
                if cells > MAX_CELLS:
                    break
                f, v = c_f.value, c_v.value
                if f is None or f == "":
                    continue
                filled += 1
                if isinstance(f, str) and f.startswith("=") and v is None:
                    blank.append((_cell(ws_f.title, c_f), f))
                if isinstance(v, str) and v.strip().upper() in ERRORS:
                    errors.append(f"{_cell(ws_f.title, c_f)} {v.strip()}")
        sheets.append(f"{ws_f.title} ({filled} filled cells)" if filled else f"{ws_f.title} (empty)")
        if not filled:
            r.problems.append(f"sheet '{ws_f.title}' is empty")
    wb_f.close()
    wb_v.close()
    r.summary = f"{len(sheets)} sheet{'s' * (len(sheets) != 1)}: " + ", ".join(sheets)
    if blank:
        where = ", ".join(f"{c} {f[:40]}" for c, f in blank[:3])
        one = len(blank) == 1
        r.problems.append(
            f"{len(blank)} formula cell{'' if one else 's'} {'has' if one else 'have'} no saved "
            f"value (e.g. {where}): "
            "a reader that doesn't recalculate (pandas, previews, graders) sees them blank. Write "
            "the computed numbers as values (the formula can go in a note or a column beside it)")
    if errors:
        r.problems.append(f"{len(errors)} cell{'s' * (len(errors) != 1)} "
                          f"show{'s' * (len(errors) == 1)} an error: " + ", ".join(errors[:5]))


def _text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


def _csv(path: Path, r: Result) -> None:
    text = _text(path)
    delim = "\t" if path.suffix.lower() == ".tsv" else ","
    if delim == "," and text.count(";") > text.count(",") * 2:
        delim = ";"
    rows = [row for row in csv.reader(io.StringIO(text), delimiter=delim) if any(c.strip() for c in row)]
    if not rows:
        r.problems.append("no rows")
        return
    header, body = rows[0], rows[1:]
    r.summary = f"{len(body)} data rows, columns: " + ", ".join(h.strip() or "(blank)" for h in header)[:300]
    if not body:
        r.problems.append("a header and no data rows")
    ragged = [i + 2 for i, row in enumerate(body) if len(row) != len(header)]
    if ragged:
        one = len(ragged) == 1
        r.problems.append(f"{len(ragged)} row{'' if one else 's'} {'has' if one else 'have'} a "
                          f"different number of fields from the header's {len(header)} (line"
                          f"{'' if one else 's'} {', '.join(map(str, ragged[:5]))}): a value with "
                          "an unquoted comma shifts every column after it")
    names = [h.strip().lower() for h in header]
    if dupes := sorted({h for h in names if h and names.count(h) > 1}):
        r.problems.append("repeated column names: " + ", ".join(dupes))
    if "" in names and len(header) > 1:
        r.problems.append("a column has no name")


def _docx(path: Path, r: Result) -> None:
    import docx

    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        parts += [c.text for row in t.rows for c in row.cells]
    text = "\n".join(parts)
    words = len(text.split())
    r.summary = f"{words:,} words, {len(d.tables)} table{'s' * (len(d.tables) != 1)}"
    if not words:
        r.problems.append("the document has no text")
    _placeholders(text, r)


def _pdf(path: Path, r: Result) -> None:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = len(reader.pages)
    r.summary = f"{pages} page{'s' * (pages != 1)}"
    if not pages:
        r.problems.append("no pages")
        return
    text = "".join((p.extract_text() or "") for p in reader.pages[:3])
    if not text.strip():
        r.summary += ", no text layer on the first pages (a scan or an image?)"
    _placeholders(text, r)


def _placeholders(text: str, r: Result) -> None:
    found = list(dict.fromkeys(m.group(0) for m in PLACEHOLDER.finditer(text)))
    if found:
        r.problems.append("placeholder text left in: " + ", ".join(repr(f[:40]) for f in found[:5]))


_LINK = re.compile(r"""(?:href|src)\s*=\s*["']([^"'#?]+)""", re.I)


def _html(path: Path, r: Result) -> None:
    text = _text(path)
    _placeholders(re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", text), r)
    missing = []
    for target in dict.fromkeys(_LINK.findall(text)):
        if re.match(r"^(?:[a-z][a-z0-9+.-]*:|//)", target, re.I) or not target.strip():
            continue
        if not (path.parent / target.strip()).exists():
            missing.append(target.strip())
    if missing:
        r.problems.append("links to files that don't exist: " + ", ".join(missing[:5]))
    r.summary = f"{len(text):,} characters"


def check_file(path: Path, rel: str) -> Result:
    r = Result(rel)
    ext = path.suffix.lower()
    try:
        size = path.stat().st_size
        if size == 0:
            r.problems.append("the file is empty (0 bytes)")
            return r
        if size > MAX_BYTES:
            r.summary = f"{size // 1_000_000} MB, too big to check"
            return r
        if ext in (".xlsx", ".xlsm"):
            _xlsx(path, r)
        elif ext in (".csv", ".tsv"):
            _csv(path, r)
        elif ext == ".docx":
            _docx(path, r)
        elif ext == ".pdf":
            _pdf(path, r)
        elif ext == ".json":
            try:
                data = json.loads(_text(path))
                r.summary = (f"a list of {len(data)}" if isinstance(data, list)
                             else f"an object with {len(data)} keys" if isinstance(data, dict)
                             else "a value")
            except ValueError as e:
                r.problems.append(f"not valid JSON: {e}")
        elif ext in (".html", ".htm"):
            _html(path, r)
        elif ext in (".md", ".txt"):
            text = _text(path)
            r.summary = f"{len(text.split()):,} words"
            _placeholders(text, r)
    except Exception as e:  # noqa: BLE001 - a file its reader can't open is the finding
        r.problems.append(f"can't be opened as {ext[1:].upper()}: {type(e).__name__}: {e}"[:300])
    return r


def _walk(root: Path, since: float | None = None, exts: set[str] = CHECKED) -> list[Path]:
    """Checkable files under root, newest first; only those changed since `since` if given."""
    from app.tools.workspace import SKIP_DIRS

    found: list[tuple[float, Path]] = []
    seen = 0
    for base, dirs, files in os.walk(root):
        dirs[:] = [d for d in dirs if d not in SKIP_DIRS and d not in SKIP]
        for f in files:
            seen += 1
            if seen > WALK_CAP:
                break
            p = Path(base) / f
            if p.suffix.lower() not in exts:
                continue
            try:
                m = p.stat().st_mtime
            except OSError:
                continue
            if since is None or m >= since:
                found.append((m, p))
        if seen > WALK_CAP:
            break
    return [p for _, p in sorted(found, key=lambda x: x[0], reverse=True)]


def check_paths(root: Path, paths: list[Path]) -> list[Result]:
    root = Path(root)
    out = []
    for p in paths[:MAX_FILES]:
        try:
            rel = p.resolve().relative_to(root.resolve()).as_posix()
        except ValueError:
            rel = p.name
        out.append(check_file(p, rel))
    return out


def check_tree(root: Path, exts: set[str] = CHECKED) -> list[Result]:
    return check_paths(root, _walk(Path(root), exts=exts))


def report(results: list[Result], *, problems_only: bool = False) -> str:
    rows = [r for r in results if r.problems] if problems_only else results
    return "\n".join(r.line(problems_only=problems_only) for r in rows)


def expected_files(request: str) -> list[str]:
    """File names the request mentions, as written: quoted ones (which may have spaces), then
    bare ones."""
    names: list[str] = []
    quoted = list(QUOTED.finditer(request))
    bare = [m for m in BARE.finditer(request)
            if not any(q.start() <= m.start() < q.end() for q in quoted)]
    for m in [*quoted, *bare]:
        name = m.group(1).strip()
        if name and Path(name).name.lower() not in (Path(n).name.lower() for n in names):
            names.append(name)
    return names


def missing_files(root: Path, names: list[str]) -> list[str]:
    """The named files found nowhere in the workspace (matched by file name, any folder)."""
    have = {p.name.lower() for p in _walk(Path(root), exts=CHECKED | {".pptx"})}
    return [n for n in names if Path(n.replace("\\", "/")).name.lower() not in have]


def changed_since(root: Path, since: float) -> list[Path]:
    return _walk(Path(root), since=since, exts=OFFICE)


def after_tool_note(root: Path, name: str, args: dict, started: float) -> str:
    """What the files a write or a command just produced look like to their reader: only
    problems, and nothing when there are none."""
    root = Path(root)
    if name in ("write_file", "edit_file"):
        rel = str(args.get("path") or "")
        p = (root / rel) if rel else None
        paths = [p] if p is not None and p.suffix.lower() in OFFICE and p.is_file() else []
    elif name == "run_command":
        paths = changed_since(root, started - 1)[:10]
    else:
        return ""
    found = [r for r in check_paths(root, paths) if r.problems]
    if not found:
        return ""
    return ("\n\n[Pantheon reopened the files this produced, as their reader will:\n"
            + report(found, problems_only=True) + "]")
