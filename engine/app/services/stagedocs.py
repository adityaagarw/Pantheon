"""Stage documents: spreadsheets, PDFs, notebooks and Markdown the user and agents share.

Every document is one row with its content as JSON and a version. Each change is
published as ``stage_doc.updated`` with what changed, so an open document applies
it live without reloading.

- sheet:    {"sheets": [{"name", "cells": {"A1": {"v": value, "f": "=formula"}}}]}
- pdf:      {"attachmentId", "name", "pages", "page", "highlights": [{"id", "page",
            "text", "note", "by"}]}
- notebook: {"cells": [{"id", "type": "code"|"markdown", "source", "outputs", "count",
            "status": "idle"|"running"|"ok"|"error", "by"}]}
- markdown: {"text"}
"""

from __future__ import annotations

import asyncio
import copy
import csv
import io
import json
import re
from collections import defaultdict
from typing import Any

from sqlalchemy import delete, select

from app.core.db import SessionLocal
from app.core.ids import new_id
from app.events import bus
from app.models import Attachment, StageDoc
from app.services import formulas
from app.services import secrets as secrets_svc

KINDS = ("sheet", "pdf", "notebook", "markdown")
MAX_CELLS = 200_000
MAX_TEXT = 2_000_000
_locks: dict[str, asyncio.Lock] = defaultdict(asyncio.Lock)


class DocError(ValueError):
    pass


def to_dict(d: StageDoc, *, full: bool = False) -> dict[str, Any]:
    out: dict[str, Any] = {
        "id": d.id, "orgId": d.org_id, "kind": d.kind, "title": d.title, "version": d.version,
        "createdBy": d.created_by, "updatedBy": d.updated_by,
        "createdAt": d.created_at.isoformat() if d.created_at else None,
        "updatedAt": d.updated_at.isoformat() if d.updated_at else None,
        "summary": summary(d),
    }
    if full:
        out["data"] = d.data
    return out


def summary(d: StageDoc) -> str:
    data = d.data or {}
    if d.kind == "sheet":
        sheets = data.get("sheets") or []
        n = sum(len(s.get("cells") or {}) for s in sheets)
        return f"{len(sheets)} sheet(s), {n} cells"
    if d.kind == "pdf":
        return f"{data.get('name', 'PDF')}, {data.get('pages') or '?'} pages"
    if d.kind == "markdown":
        text = data.get("text") or ""
        return f"{len(text.split())} words, {text.count(chr(10)) + 1 if text else 0} lines"
    cells = data.get("cells") or []
    return f"{len(cells)} cells"


async def _publish(d: StageDoc, **extra: Any) -> None:
    # Live only: documents load fresh, and outputs (charts) would bloat the event log.
    await bus.publish("stage_doc.updated", {**to_dict(d), **extra}, org_id=d.org_id,
                      persist=False)


async def create(org_id: str, kind: str, title: str, data: dict[str, Any] | None = None, *,
                 created_by: str = "user", by_name: str = "user") -> StageDoc:
    if kind not in KINDS:
        raise DocError(f"kind must be one of {', '.join(KINDS)}")
    if data is None:
        data = {"sheet": {"sheets": [{"name": "Sheet1", "cells": {}}]},
                "notebook": {"cells": [_new_cell("code", "", by_name)]},
                "markdown": {"text": ""},
                "pdf": {}}[kind]
    d = StageDoc(org_id=org_id, kind=kind, title=(title.strip() or kind.title())[:200], data=data,
                 created_by=created_by, updated_by=by_name)
    async with SessionLocal() as session:
        session.add(d)
        await session.commit()
    await _publish(d)
    return d


async def get(doc_id: str) -> StageDoc:
    async with SessionLocal() as session:
        d = await session.get(StageDoc, doc_id)
    if d is None:
        raise DocError("document not found")
    return d


async def for_org(org_id: str, kind: str | None = None) -> list[StageDoc]:
    async with SessionLocal() as session:
        q = select(StageDoc).where(StageDoc.org_id == org_id)
        if kind:
            q = q.where(StageDoc.kind == kind)
        return list((await session.execute(q.order_by(StageDoc.updated_at.desc()))).scalars())


async def resolve(org_id: str, ref: str, kind: str | None = None) -> StageDoc:
    """A document by id or (case-insensitive) title."""
    docs = await for_org(org_id, kind)
    ref = (ref or "").strip()
    for d in docs:
        if d.id == ref:
            return d
    for d in docs:
        if d.title.lower() == ref.lower():
            return d
    what = f"{kind} " if kind else ""
    names = ", ".join(f"'{d.title}'" for d in docs[:10]) or "none yet"
    raise DocError(f"no {what}document '{ref}' (there are: {names})")


async def _save(doc_id: str, mutate, by_name: str) -> tuple[StageDoc, Any]:
    """Apply ``mutate(data) -> result`` under the document's lock and bump its version."""
    async with _locks[doc_id]:
        async with SessionLocal() as session:
            d = await session.get(StageDoc, doc_id)
            if d is None:
                raise DocError("document not found")
            data = copy.deepcopy(d.data or {})
            result = mutate(data)
            d.data = data
            d.version += 1
            d.updated_by = by_name
            await session.commit()
            await session.refresh(d)
    return d, result


async def rename(doc_id: str, title: str, by_name: str = "user") -> StageDoc:
    if not title.strip():
        raise DocError("a title is required")
    async with SessionLocal() as session:
        d = await session.get(StageDoc, doc_id)
        if d is None:
            raise DocError("document not found")
        d.title, d.updated_by = title.strip()[:200], by_name
        d.version += 1
        await session.commit()
    await _publish(d)
    return d


async def delete_doc(doc_id: str) -> None:
    d = await get(doc_id)
    async with SessionLocal() as session:
        await session.execute(delete(StageDoc).where(StageDoc.id == doc_id))
        await session.commit()
    if d.kind == "notebook":
        from app.services import notebook_kernel

        try:
            await notebook_kernel.restart(doc_id)
        except Exception:  # noqa: BLE001 - the sandbox may be off; nothing to clean then
            pass
    await bus.publish("stage_doc.deleted", {"id": doc_id}, org_id=d.org_id)


# --- spreadsheets ---------------------------------------------------------------------------


def _book(data: dict[str, Any]) -> dict[str, dict[str, dict[str, Any]]]:
    return {s["name"]: s.setdefault("cells", {}) for s in data.setdefault("sheets", [])}


def _sheet(data: dict[str, Any], name: str | None) -> dict[str, Any]:
    sheets = data.setdefault("sheets", [])
    if not sheets:
        sheets.append({"name": "Sheet1", "cells": {}})
    if not name:
        return sheets[0]
    for s in sheets:
        if s["name"].lower() == name.lower():
            return s
    sheets.append({"name": name[:60], "cells": {}})  # writing to a new sheet creates it
    return sheets[-1]


def _coerce(value: Any) -> Any:
    """Numbers stay numbers; numeric text becomes a number, like typing it into a cell."""
    if isinstance(value, bool) or value is None:
        return value
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value)
    if re.fullmatch(r"\s*-?\d+(?:\.\d+)?(?:[eE][+-]?\d+)?\s*", text):
        return float(text)
    return text


def _cell_from(value: Any = None, formula: str | None = None) -> dict[str, Any] | None:
    if formula:
        return {"f": formula if formula.startswith("=") else f"={formula}"}
    if isinstance(value, str) and value.startswith("=") and len(value) > 1:
        return {"f": value}
    if value is None or value == "":
        return None
    return {"v": _coerce(value)}


def _apply_cells(data: dict[str, Any], sheet: str | None,
                 updates: dict[str, dict[str, Any] | None]) -> dict[str, dict[str, Any] | None]:
    """Set (or clear, with None) cells, recalculate, and return what changed."""
    target = _sheet(data, sheet)
    cells = target["cells"]
    for ref, cell in updates.items():
        ref = ref.upper().replace("$", "")
        formulas.parse_cell(ref)  # validates
        if cell is None:
            cells.pop(ref, None)
        else:
            cells[ref] = cell
    if sum(len(c) for c in _book(data).values()) > MAX_CELLS:
        raise DocError(f"a spreadsheet can hold at most {MAX_CELLS} cells")
    formulas.recalculate(_book(data))
    # Changed cells plus formulas whose values may have moved with them.
    changed = {ref.upper().replace("$", ""): cells.get(ref.upper().replace("$", ""))
               for ref in updates}
    for ref, cell in cells.items():
        if cell.get("f"):
            changed[ref] = cell
    return {target["name"]: changed}


async def write_cells(doc_id: str, writes: list[dict[str, Any]], by_name: str) -> dict[str, Any]:
    """writes: [{"cell": "B2", "value"?: .., "formula"?: "=..", "sheet"?: ..}]"""
    def mutate(data: dict[str, Any]) -> dict[str, Any]:
        by_sheet: dict[str | None, dict[str, Any]] = defaultdict(dict)
        for w in writes:
            by_sheet[w.get("sheet")][str(w["cell"])] = _cell_from(w.get("value"), w.get("formula"))
        changes: dict[str, Any] = {}
        for sheet, updates in by_sheet.items():
            for name, cells in _apply_cells(data, sheet, updates).items():
                changes.setdefault(name, {}).update(cells)
        return changes

    try:
        d, changes = await _save(doc_id, mutate, by_name)
    except ValueError as e:
        raise DocError(str(e)) from None
    await _publish(d, changes=changes)
    return changes


async def write_range(doc_id: str, start: str, rows: list[list[Any]], by_name: str,
                      sheet: str | None = None) -> dict[str, Any]:
    r0, c0 = formulas.parse_cell(start)
    writes = [{"cell": formulas.cell_name(r0 + i, c0 + j), "value": v, "sheet": sheet}
              for i, row in enumerate(rows) for j, v in enumerate(row)]
    return await write_cells(doc_id, writes, by_name)


async def apply_user_changes(doc_id: str, changes: dict[str, dict[str, Any]],
                             sheet_names: list[str] | None = None) -> StageDoc:
    """The open spreadsheet saves what the user changed: {sheet: {ref: cell | None}}."""
    def mutate(data: dict[str, Any]) -> None:
        if sheet_names:
            existing = {s["name"]: s for s in data.get("sheets", [])}
            data["sheets"] = [existing.get(n) or {"name": n, "cells": {}} for n in sheet_names]
        for sheet, updates in changes.items():
            target = _sheet(data, sheet)
            for ref, cell in updates.items():
                ref = ref.upper()
                if cell is None or (cell.get("v") in (None, "") and not cell.get("f")):
                    target["cells"].pop(ref, None)
                else:
                    keep = {k: cell[k] for k in ("v", "f") if cell.get(k) not in (None, "")}
                    if "v" in keep and not keep.get("f"):
                        keep["v"] = _coerce(keep["v"])  # the grid sends typed numbers as text
                    target["cells"][ref] = keep
        formulas.recalculate(_book(data))

    d, _ = await _save(doc_id, mutate, "user")
    await _publish(d, changes=changes, source="user")
    return d


def used_range(cells: dict[str, Any]) -> tuple[int, int, int, int] | None:
    if not cells:
        return None
    rc = [formulas.parse_cell(ref) for ref in cells]
    return 0, 0, max(r for r, _ in rc), max(c for _, c in rc)


def read_sheet(d: StageDoc, sheet: str | None = None, rng: str | None = None,
               max_rows: int = 200, max_cols: int = 26) -> str:
    data = d.data or {}
    sheets = data.get("sheets") or []
    if not sheets:
        return f"'{d.title}' is empty."
    target = sheets[0]
    if sheet:
        target = next((s for s in sheets if s["name"].lower() == sheet.lower()), None)
        if target is None:
            raise DocError(f"no sheet '{sheet}' (sheets: {', '.join(s['name'] for s in sheets)})")
    cells = target.get("cells") or {}
    bounds = formulas.parse_range(rng) if rng else used_range(cells)
    header = (f"Spreadsheet '{d.title}' ({d.id}), sheet '{target['name']}'"
              + (f" of {', '.join(s['name'] for s in sheets)}" if len(sheets) > 1 else ""))
    if bounds is None:
        return header + ": empty."
    r1, c1, r2, c2 = bounds
    clipped = r2 - r1 + 1 > max_rows or c2 - c1 + 1 > max_cols
    r2, c2 = min(r2, r1 + max_rows - 1), min(c2, c1 + max_cols - 1)
    lines = [header + f", {formulas.cell_name(r1, c1)}:{formulas.cell_name(r2, c2)}:", "",
             "|   | " + " | ".join(formulas.col_letters(c) for c in range(c1, c2 + 1)) + " |",
             "|---|" + "---|" * (c2 - c1 + 1)]
    with_formulas = []
    for r in range(r1, r2 + 1):
        row = []
        for c in range(c1, c2 + 1):
            ref = formulas.cell_name(r, c)
            cell = cells.get(ref) or {}
            row.append(formulas.display(cell.get("v")).replace("|", "\\|").replace("\n", " "))
            if cell.get("f"):
                with_formulas.append(f"{ref} {cell['f']} → {formulas.display(cell.get('v'))}")
        lines.append(f"| {r + 1} | " + " | ".join(row) + " |")
    if with_formulas:
        lines += ["", "Formulas:"] + with_formulas[:200]
    if clipped:
        lines.append("\n(more cells outside this view: give a range)")
    return "\n".join(lines)


def sheets_from_file(name: str, raw: bytes) -> list[dict[str, Any]]:
    """A spreadsheet's sheets from an uploaded .xlsx or .csv (formulas kept)."""
    if name.lower().endswith((".csv", ".tsv", ".txt")):
        text = raw.decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:4096]) if text.strip() else csv.excel
        cells: dict[str, Any] = {}
        for r, row in enumerate(csv.reader(io.StringIO(text), dialect)):
            for c, v in enumerate(row):
                cell = _cell_from(v)
                if cell:
                    cells[formulas.cell_name(r, c)] = cell
        return [{"name": "Sheet1", "cells": cells}]
    if name.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb_f = load_workbook(io.BytesIO(raw), data_only=False, read_only=True)
        wb_v = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
        out = []
        for ws in wb_f.worksheets:
            vals = wb_v[ws.title]
            cells = {}
            for row, vrow in zip(ws.iter_rows(), vals.iter_rows(), strict=False):
                for cell, vcell in zip(row, vrow, strict=False):
                    if cell.value is None or not hasattr(cell, "coordinate"):
                        continue
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        cells[cell.coordinate] = {"f": cell.value, "v": _jsonable(vcell.value)}
                    else:
                        cells[cell.coordinate] = {"v": _jsonable(cell.value)}
            out.append({"name": ws.title[:60], "cells": cells})
        # Files saved by programs carry no cached results; Excel's own are kept for
        # functions the server doesn't know.
        formulas.recalculate({s["name"]: s["cells"] for s in out})
        return out or [{"name": "Sheet1", "cells": {}}]
    raise DocError("import a .xlsx or .csv file")


def _jsonable(v: Any) -> Any:
    if v is None or isinstance(v, (bool, str)):
        return v
    if isinstance(v, (int, float)):
        return float(v)
    return str(v)  # dates and times as text


def sheet_xlsx(d: StageDoc) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for s in (d.data or {}).get("sheets") or [{"name": "Sheet1", "cells": {}}]:
        ws = wb.create_sheet(s["name"][:31] or "Sheet")
        for ref, cell in (s.get("cells") or {}).items():
            v = cell.get("f") or cell.get("v")
            ws[ref] = int(v) if isinstance(v, float) and v.is_integer() else v
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- PDFs -----------------------------------------------------------------------------------


async def open_pdf(org_id: str, attachment: Attachment, title: str | None, *,
                   created_by: str, by_name: str) -> StageDoc:
    if attachment.kind != "pdf" and not attachment.name.lower().endswith(".pdf"):
        raise DocError(f"'{attachment.name}' isn't a PDF")
    pages = (attachment.info or {}).get("pages")
    return await create(org_id, "pdf", title or attachment.name, {
        "attachmentId": attachment.id, "name": attachment.name, "pages": pages, "page": 1,
        "highlights": []}, created_by=created_by, by_name=by_name)


async def show_pdf(doc_id: str, page: int | None, highlight: str | None, note: str | None,
                   by_name: str, clear: bool = False) -> StageDoc:
    def mutate(data: dict[str, Any]) -> None:
        hl = [] if clear else data.setdefault("highlights", [])
        if page:
            pages = data.get("pages")
            if page < 1 or (pages and page > pages):
                raise DocError(f"page must be between 1 and {pages or '…'}")
            data["page"] = page
        if highlight:
            hl.append({"id": new_id("hl"), "page": page or data.get("page") or 1,
                       "text": highlight.strip()[:500], "note": (note or "")[:500], "by": by_name})
        data["highlights"] = hl[-100:]

    d, _ = await _save(doc_id, mutate, by_name)
    await _publish(d, focus={"page": (d.data or {}).get("page", 1)}, data=d.data)
    return d


def pdf_page_text(path: str, page: int) -> tuple[str, int]:
    from pypdf import PdfReader

    reader = PdfReader(path)
    total = len(reader.pages)
    if not 1 <= page <= total:
        raise DocError(f"page must be between 1 and {total}")
    return (reader.pages[page - 1].extract_text() or "").strip(), total


# --- notebooks ------------------------------------------------------------------------------


def _new_cell(kind: str, source: str, by: str) -> dict[str, Any]:
    return {"id": new_id("cell"), "type": kind if kind in ("code", "markdown") else "code",
            "source": source, "outputs": [], "count": None, "status": "idle", "by": by}


def _find(data: dict[str, Any], ref: str | int) -> int:
    cells = data.get("cells") or []
    if isinstance(ref, int) or (isinstance(ref, str) and ref.isdigit()):
        i = int(ref) - 1  # 1-based for people and agents
        if 0 <= i < len(cells):
            return i
    for i, c in enumerate(cells):
        if c["id"] == ref:
            return i
    raise DocError(f"no cell '{ref}' (cells are numbered 1-{len(cells)})")


async def add_cell(doc_id: str, source: str, kind: str, by_name: str,
                   after: str | int | None = None) -> dict[str, Any]:
    cell = _new_cell(kind, secrets_svc.redact((await get(doc_id)).org_id, source), by_name)

    def mutate(data: dict[str, Any]) -> None:
        cells = data.setdefault("cells", [])
        # An untouched empty first cell is replaced rather than kept above.
        if len(cells) == 1 and not cells[0]["source"].strip() and not cells[0]["outputs"] \
                and after is None:
            cells.clear()
        idx = len(cells) if after is None else _find(data, after) + 1
        cells.insert(idx, cell)

    d, _ = await _save(doc_id, mutate, by_name)
    await _publish(d, data=d.data)
    return cell


async def edit_cell(doc_id: str, ref: str | int, *, source: str | None = None,
                    kind: str | None = None, by_name: str) -> dict[str, Any]:
    org_id = (await get(doc_id)).org_id

    def mutate(data: dict[str, Any]) -> dict[str, Any]:
        c = data["cells"][_find(data, ref)]
        if source is not None:
            c["source"] = secrets_svc.redact(org_id, source)
        if kind in ("code", "markdown"):
            c["type"] = kind
        c["by"] = by_name
        return c

    d, cell = await _save(doc_id, mutate, by_name)
    await _publish(d, data=d.data)
    return cell


async def delete_cell(doc_id: str, ref: str | int, by_name: str) -> None:
    def mutate(data: dict[str, Any]) -> None:
        data["cells"].pop(_find(data, ref))

    d, _ = await _save(doc_id, mutate, by_name)
    await _publish(d, data=d.data)


async def move_cell(doc_id: str, ref: str | int, to: int, by_name: str) -> None:
    def mutate(data: dict[str, Any]) -> None:
        cells = data["cells"]
        cell = cells.pop(_find(data, ref))
        cells.insert(max(0, min(to - 1, len(cells))), cell)

    d, _ = await _save(doc_id, mutate, by_name)
    await _publish(d, data=d.data)


async def run_cell(doc_id: str, ref: str | int, by_name: str,
                   timeout: float = 300) -> dict[str, Any]:
    from app.services import notebook_kernel

    def start(data: dict[str, Any]) -> dict[str, Any]:
        c = data["cells"][_find(data, ref)]
        c["status"] = "running" if c["type"] == "code" else "idle"
        return dict(c)

    d, cell = await _save(doc_id, start, by_name)
    if cell["type"] != "code":
        return cell
    from app.core import guard

    if guard.touches_data(cell["source"]):
        result = {"count": None, "status": "error", "outputs": [
            {"type": "error", "ename": "PermissionError", "evalue": guard.OFF_LIMITS,
             "traceback": ""}]}

        def refuse(data: dict[str, Any]) -> dict[str, Any]:
            c = data["cells"][_find(data, ref)]
            c.update(status="error", outputs=result["outputs"])
            return dict(c)

        d, cell = await _save(doc_id, refuse, by_name)
        await _publish(d, data=d.data)
        return cell
    await _publish(d, data=d.data)
    try:
        result = await notebook_kernel.execute(doc_id, cell["source"], timeout)
    except notebook_kernel.KernelError as e:
        result = {"count": None, "status": "error", "outputs": [
            {"type": "error", "ename": "KernelUnavailable", "evalue": str(e), "traceback": ""}]}
    org_id = d.org_id

    def finish(data: dict[str, Any]) -> dict[str, Any]:
        try:
            c = data["cells"][_find(data, cell["id"])]
        except DocError:
            return cell  # deleted while running
        c["outputs"] = _clean_outputs(org_id, result.get("outputs") or [])
        c["count"] = result.get("count")
        c["status"] = result.get("status", "ok")
        return c

    d, done = await _save(doc_id, finish, by_name)
    await _publish(d, data=d.data)
    return done


def _clean_outputs(org_id: str, outputs: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out = []
    for o in outputs[:50]:
        o = dict(o)
        for k in ("text", "evalue", "traceback", "html"):
            if isinstance(o.get(k), str):
                o[k] = secrets_svc.redact(org_id, o[k])
        out.append(o)
    return out


async def run_all(doc_id: str, by_name: str) -> list[dict[str, Any]]:
    d = await get(doc_id)
    results = []
    for c in (d.data or {}).get("cells") or []:
        if c["type"] != "code" or not c["source"].strip():
            continue
        done = await run_cell(doc_id, c["id"], by_name)
        results.append(done)
        if done.get("status") == "error":
            break
    return results


def read_notebook(d: StageDoc, max_output: int = 1500) -> str:
    cells = (d.data or {}).get("cells") or []
    lines = [f"Notebook '{d.title}' ({d.id}), {len(cells)} cells:"]
    for i, c in enumerate(cells, 1):
        mark = f"[{c['count']}]" if c.get("count") else "[ ]"
        head = f"\n--- cell {i} ({c['type']}{', ' + mark if c['type'] == 'code' else ''}"
        head += f", {c['status']})" if c.get("status") not in (None, "idle", "ok") else ")"
        lines.append(head)
        lines.append(c["source"] or "(empty)")
        for o in c.get("outputs") or []:
            lines.append(output_text(o, max_output))
    return "\n".join(lines)


def output_text(o: dict[str, Any], cap: int = 1500) -> str:
    t = o.get("type")
    if t in ("stream", "result"):
        text = o.get("text", "")
    elif t == "image":
        text = "[a chart/image]"
    elif t == "error":
        text = f"{o.get('ename')}: {o.get('evalue')}\n{o.get('traceback', '')}".strip()
    else:
        text = json.dumps(o)[:200]
    if len(text) > cap:
        text = text[: cap // 2] + "\n…\n" + text[-cap // 2:]
    return "> " + text.replace("\n", "\n> ")


def notebook_from_ipynb(raw: bytes) -> dict[str, Any]:
    try:
        nb = json.loads(raw.decode("utf-8"))
    except (ValueError, UnicodeDecodeError):
        raise DocError("that isn't a Jupyter notebook (.ipynb)") from None
    cells = []
    for c in nb.get("cells") or []:
        src = c.get("source") or ""
        src = "".join(src) if isinstance(src, list) else str(src)
        kind = "markdown" if c.get("cell_type") == "markdown" else "code"
        cell = _new_cell(kind, src, "user")
        if kind == "code":
            cell["count"] = c.get("execution_count")
            cell["outputs"] = [o for o in (_from_nb_output(x) for x in c.get("outputs") or []) if o]
        cells.append(cell)
    return {"cells": cells or [_new_cell("code", "", "user")]}


def _from_nb_output(o: dict[str, Any]) -> dict[str, Any] | None:
    def text(v: Any) -> str:
        return "".join(v) if isinstance(v, list) else str(v or "")

    kind = o.get("output_type")
    if kind == "stream":
        return {"type": "stream", "name": o.get("name", "stdout"), "text": text(o.get("text"))}
    if kind == "error":
        return {"type": "error", "ename": o.get("ename", ""), "evalue": o.get("evalue", ""),
                "traceback": re.sub(r"\x1b\[[0-9;]*m", "", "\n".join(o.get("traceback") or []))}
    if kind in ("execute_result", "display_data"):
        data = o.get("data") or {}
        if "image/png" in data:
            return {"type": "image", "data": text(data["image/png"]).replace("\n", "")}
        out = {"type": "result", "text": text(data.get("text/plain"))}
        if "text/html" in data:
            out["html"] = text(data["text/html"])
        return out
    return None


def notebook_ipynb(d: StageDoc) -> bytes:
    cells = []
    for c in (d.data or {}).get("cells") or []:
        if c["type"] == "markdown":
            cells.append({"cell_type": "markdown", "metadata": {},
                          "source": c["source"].splitlines(keepends=True)})
            continue
        outputs = []
        for o in c.get("outputs") or []:
            if o["type"] == "stream":
                outputs.append({"output_type": "stream", "name": o.get("name", "stdout"),
                                "text": o.get("text", "").splitlines(keepends=True)})
            elif o["type"] == "error":
                outputs.append({"output_type": "error", "ename": o.get("ename", ""),
                                "evalue": o.get("evalue", ""),
                                "traceback": o.get("traceback", "").splitlines()})
            elif o["type"] == "image":
                outputs.append({"output_type": "display_data", "metadata": {},
                                "data": {"image/png": o.get("data", "")}})
            elif o["type"] == "result":
                data = {"text/plain": o.get("text", "")}
                if o.get("html"):
                    data["text/html"] = o["html"]
                outputs.append({"output_type": "execute_result", "metadata": {}, "data": data,
                                "execution_count": c.get("count")})
        cells.append({"cell_type": "code", "metadata": {}, "execution_count": c.get("count"),
                      "source": c["source"].splitlines(keepends=True), "outputs": outputs})
    nb = {"nbformat": 4, "nbformat_minor": 5, "cells": cells,
          "metadata": {"kernelspec": {"name": "python3", "display_name": "Python 3",
                                      "language": "python"},
                       "language_info": {"name": "python"}}}
    return json.dumps(nb, indent=1).encode()


# --- Markdown -------------------------------------------------------------------------------


class Conflict(DocError):
    """The text changed since the version an edit was based on."""


def _text(org_id: str, text: str) -> str:
    if len(text) > MAX_TEXT:
        raise DocError(f"a Markdown document can hold at most {MAX_TEXT} characters")
    return secrets_svc.redact(org_id, text)


async def set_text(doc_id: str, text: str, by_name: str, base: int | None = None) -> StageDoc:
    """Replace the whole text. ``base`` is the version the editor started from: if someone
    else changed it since, nothing is saved and Conflict is raised."""
    org_id = (await get(doc_id)).org_id
    clean = _text(org_id, text)
    async with _locks[doc_id]:
        async with SessionLocal() as session:
            d = await session.get(StageDoc, doc_id)
            if d is None:
                raise DocError("document not found")
            if base is not None and base != d.version:
                raise Conflict(f"'{d.title}' was changed by {d.updated_by} meanwhile")
            d.data = {**(d.data or {}), "text": clean}
            d.version += 1
            d.updated_by = by_name
            await session.commit()
            await session.refresh(d)
    await _publish(d, data=d.data, source=by_name)
    return d


async def edit_text(doc_id: str, by_name: str, *, find: str | None = None,
                    replace: str = "", append: str | None = None) -> StageDoc:
    """An agent's change: replace one exact passage, or add to the end."""
    org_id = (await get(doc_id)).org_id

    def mutate(data: dict[str, Any]) -> None:
        text = data.get("text") or ""
        if find:
            n = text.count(find)
            if n == 0:
                raise DocError("that passage isn't in the document (copy it exactly from md_read)")
            if n > 1:
                raise DocError(f"that passage appears {n} times: include more around it")
            text = text.replace(find, replace, 1)
        if append:
            text = (text.rstrip("\n") + "\n\n" if text.strip() else "") + append
        data["text"] = _text(org_id, text)

    d, _ = await _save(doc_id, mutate, by_name)
    await _publish(d, data=d.data, source=by_name)
    return d


def read_markdown(d: StageDoc, start: int = 1, lines: int = 400) -> str:
    rows = ((d.data or {}).get("text") or "").splitlines()
    if not rows:
        return f"Markdown '{d.title}' ({d.id}) is empty."
    start = max(1, start)
    end = min(len(rows), start + lines - 1)
    body = "\n".join(rows[start - 1:end])
    head = f"Markdown '{d.title}' ({d.id}), lines {start}-{end} of {len(rows)}:"
    more = f"\n\n(more below: start={end + 1})" if end < len(rows) else ""
    return f"{head}\n\n{body}{more}"


def markdown_from_file(raw: bytes) -> dict[str, Any]:
    return {"text": raw.decode("utf-8-sig", errors="replace")[:MAX_TEXT]}
